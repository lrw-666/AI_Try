from __future__ import annotations

import argparse
import importlib.util
import json
import re
import sys
import zipfile
from pathlib import Path
from types import ModuleType

from docx import Document


ALLOWED_STATUS = {"peer-reviewed", "accepted", "preprint", "status-unverified"}
ALLOWED_PRIORITY = {"core", "targeted", "reference"}
ALLOWED_READING_MODE = {"deep", "selective", "lookup"}
REQUIRED_SOURCE_FIELDS = {
    "id",
    "title",
    "authors",
    "year",
    "status",
    "url",
    "layer",
    "priority",
    "reading_mode",
    "purpose",
}
FORBIDDEN_TEXT = {
    "Lorem",
    "TBD",
    "TODO",
    "[Author]",
    "[Month YYYY]",
    "turn0search",
    "turn0view",
    "cite",
}


def validate_sources(path: Path) -> list[str]:
    """Return source-catalog errors without mutating the catalog."""
    errors: list[str] = []
    try:
        records = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"cannot read source catalog: {exc}"]

    if not isinstance(records, list) or not records:
        return ["source catalog must be a non-empty JSON array"]

    seen_ids: set[str] = set()
    seen_urls: set[str] = set()
    for index, record in enumerate(records, start=1):
        label = f"source[{index}]"
        if not isinstance(record, dict):
            errors.append(f"{label} must be an object")
            continue
        missing = REQUIRED_SOURCE_FIELDS - record.keys()
        if missing:
            errors.append(f"{label} missing fields: {sorted(missing)}")
        source_id = str(record.get("id", "")).strip()
        url = str(record.get("url", "")).strip()
        if not source_id:
            errors.append(f"{label} has an empty id")
        elif source_id in seen_ids:
            errors.append(f"duplicate source id: {source_id}")
        seen_ids.add(source_id)
        if not url.startswith("https://"):
            errors.append(f"{label} URL must use HTTPS: {url}")
        elif url in seen_urls:
            errors.append(f"duplicate source URL: {url}")
        seen_urls.add(url)
        if record.get("status") not in ALLOWED_STATUS:
            errors.append(f"{label} has unsupported status: {record.get('status')}")
        if record.get("priority") not in ALLOWED_PRIORITY:
            errors.append(f"{label} has unsupported priority: {record.get('priority')}")
        if record.get("reading_mode") not in ALLOWED_READING_MODE:
            errors.append(
                f"{label} has unsupported reading_mode: {record.get('reading_mode')}"
            )
        if not isinstance(record.get("year"), int):
            errors.append(f"{label} year must be an integer")
        for field in ("title", "authors", "layer", "purpose"):
            if not str(record.get(field, "")).strip():
                errors.append(f"{label} has empty {field}")

    if len(records) < 25:
        errors.append("source catalog must include at least 25 survey and method papers")
    deep_count = sum(item.get("reading_mode") == "deep" for item in records)
    selective_count = sum(item.get("reading_mode") == "selective" for item in records)
    if deep_count < 10:
        errors.append("source catalog must include at least 10 deep-reading papers")
    if selective_count < 10:
        errors.append("source catalog must include at least 10 selective-reading papers")
    return errors


def load_module(path: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location("report_content", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot import content module: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def flatten_text(value: object) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return "\n".join(flatten_text(item) for item in value.values())
    if isinstance(value, (list, tuple, set)):
        return "\n".join(flatten_text(item) for item in value)
    return str(value)


def validate_content(path: Path) -> list[str]:
    errors: list[str] = []
    try:
        module = load_module(path)
    except Exception as exc:  # pragma: no cover - reports import failures directly
        return [f"cannot import report content: {exc}"]

    required_names = {
        "REPORT_META",
        "REQUIRED_SECTIONS",
        "SECTIONS",
        "WEEKLY_PLAN",
        "FINALIZATION_DAY",
        "TOPIC_CRITERIA",
        "TOPIC_FAMILIES",
        "HARD_GATES",
        "RISKS",
        "APPENDICES",
    }
    missing_names = [name for name in required_names if not hasattr(module, name)]
    if missing_names:
        return [f"content module missing names: {missing_names}"]

    meta = module.REPORT_META
    if meta.get("start_date") != "2026-07-14" or meta.get("end_date") != "2026-08-31":
        errors.append("report dates must be exactly 2026-07-14 through 2026-08-31")

    missing_sections = set(module.REQUIRED_SECTIONS) - set(module.SECTIONS)
    if missing_sections:
        errors.append(f"missing required sections: {sorted(missing_sections)}")

    expected_ranges = [
        ("2026-07-14", "2026-07-19"),
        ("2026-07-20", "2026-07-26"),
        ("2026-07-27", "2026-08-02"),
        ("2026-08-03", "2026-08-09"),
        ("2026-08-10", "2026-08-16"),
        ("2026-08-17", "2026-08-23"),
        ("2026-08-24", "2026-08-30"),
    ]
    if len(module.WEEKLY_PLAN) != 7:
        errors.append("weekly plan must contain exactly seven weekly records")
    weekly_fields = {
        "start",
        "end",
        "theme",
        "hours",
        "deep_reading",
        "selective_reading",
        "concept_questions",
        "practice",
        "deliverables",
        "topic_funnel_action",
        "completion_gate",
    }
    for index, expected in enumerate(expected_ranges):
        if index >= len(module.WEEKLY_PLAN):
            break
        week = module.WEEKLY_PLAN[index]
        missing = weekly_fields - week.keys()
        if missing:
            errors.append(f"week {index + 1} missing fields: {sorted(missing)}")
        if (week.get("start"), week.get("end")) != expected:
            errors.append(f"week {index + 1} has incorrect date range")
        for field in (
            "deep_reading",
            "selective_reading",
            "concept_questions",
            "practice",
            "deliverables",
        ):
            if not week.get(field):
                errors.append(f"week {index + 1} has no {field}")
        if not week.get("completion_gate"):
            errors.append(f"week {index + 1} has no completion gate")

    final_day = module.FINALIZATION_DAY
    if final_day.get("date") != "2026-08-31":
        errors.append("finalization record must use 2026-08-31")
    if not final_day.get("deliverables"):
        errors.append("finalization record must include deliverables")

    criteria_total = sum(weight for _, _, weight, _ in module.TOPIC_CRITERIA)
    if criteria_total != 100:
        errors.append(f"topic criteria must total 100 points, got {criteria_total}")
    if len(module.TOPIC_FAMILIES) < 5:
        errors.append("at least five candidate topic families are required")
    if len(module.HARD_GATES) < 4:
        errors.append("at least four topic hard gates are required")

    report_text = flatten_text(
        [
            module.REPORT_META,
            module.SECTIONS,
            module.WEEKLY_PLAN,
            module.FINALIZATION_DAY,
            module.TOPIC_CRITERIA,
            module.TOPIC_FAMILIES,
            module.HARD_GATES,
            module.RISKS,
            module.APPENDICES,
        ]
    )
    for forbidden in FORBIDDEN_TEXT:
        if forbidden.lower() in report_text.lower():
            errors.append(f"content contains forbidden placeholder/token: {forbidden}")
    if len(report_text) < 12000:
        errors.append("report content is too short for the approved comprehensive scope")
    return errors


def _document_text(doc: Document) -> str:
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(chunks)


def validate_docx(path: Path) -> list[str]:
    errors: list[str] = []
    try:
        doc = Document(path)
    except Exception as exc:
        return [f"cannot open DOCX: {exc}"]

    if len(doc.sections) != 2:
        errors.append(f"DOCX must preserve two sections, got {len(doc.sections)}")
    if len(doc.inline_shapes) < 1:
        errors.append("DOCX must preserve the template cover image")

    text = _document_text(doc)
    required_phrases = [
        "Agent 综述学习与两年研究选题设计报告",
        "2026 年 7 月 14 日至 8 月 31 日",
        "执行摘要",
        "Agent 知识架构",
        "七周强化计划",
        "选题漏斗",
        "两年研究路线图",
        "个人研究原型成本",
        "附录",
    ]
    compact_text = re.sub(r"\s+", "", text)
    for phrase in required_phrases:
        if re.sub(r"\s+", "", phrase) not in compact_text:
            errors.append(f"DOCX missing required phrase: {phrase}")
    for forbidden in FORBIDDEN_TEXT:
        if forbidden.lower() in text.lower():
            errors.append(f"DOCX contains forbidden placeholder/token: {forbidden}")

    heading_1 = sum(paragraph.style.name == "Heading 1" for paragraph in doc.paragraphs)
    heading_2 = sum(paragraph.style.name == "Heading 2" for paragraph in doc.paragraphs)
    if heading_1 < 8:
        errors.append(f"DOCX needs at least eight Heading 1 paragraphs, got {heading_1}")
    if heading_2 < 20:
        errors.append(f"DOCX needs at least twenty Heading 2 paragraphs, got {heading_2}")
    if len(doc.tables) < 8:
        errors.append(f"DOCX needs at least eight structured tables, got {len(doc.tables)}")

    with zipfile.ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
        footer_xml = "\n".join(
            package.read(name).decode("utf-8")
            for name in package.namelist()
            if re.fullmatch(r"word/footer\d+\.xml", name)
        )
        if not re.search(r"<w:instrText[^>]*>\s*TOC\b", document_xml):
            errors.append("DOCX must preserve a TOC field")
        if "PAGE" not in footer_xml:
            errors.append("DOCX must preserve a PAGE field")
    return errors


def report_result(label: str, errors: list[str]) -> bool:
    if errors:
        print(f"FAIL: {label}")
        for error in errors:
            print(f"  - {error}")
        return False
    print(f"PASS: {label}")
    return True


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sources", type=Path)
    parser.add_argument("--content", type=Path)
    parser.add_argument("--docx", type=Path)
    args = parser.parse_args()
    if not any((args.sources, args.content, args.docx)):
        parser.error("provide at least one of --sources, --content, or --docx")

    passed = True
    if args.sources:
        passed &= report_result(
            "source catalog is structurally valid", validate_sources(args.sources)
        )
    if args.content:
        passed &= report_result(
            "report content satisfies the approved specification",
            validate_content(args.content),
        )
    if args.docx:
        passed &= report_result(
            "DOCX structure and required content are valid", validate_docx(args.docx)
        )
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
