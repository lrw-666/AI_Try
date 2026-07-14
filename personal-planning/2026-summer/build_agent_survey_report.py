from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt

from report_content import (
    APPENDICES,
    COST_ITEMS,
    COST_SUMMARY,
    FINALIZATION_DAY,
    GPU_ROLE_MATRIX,
    HARD_GATES,
    KNOWLEDGE_LAYERS,
    REPORT_META,
    RISKS,
    SECTIONS,
    TOPIC_CRITERIA,
    TOPIC_FAMILIES,
    TWO_YEAR_PHASES,
    WEEKLY_PLAN,
)


REFERENCE = Path(
    "C:/Users/10746/.codex/plugins/cache/openai-curated-remote/"
    "openai-templates/0.1.0/skills/artifact-template-design-report/"
    "assets/reference.docx"
)
EXPECTED_SHA256 = "ba1e11258ff52659318a321462a5e598c7bed33cf991329eb91150788fcd1a7b"
DEFAULT_OUTPUT = Path(__file__).with_name("agent-survey-and-research-topic-report.docx")
DEFAULT_SOURCES = Path(__file__).with_name("agent-survey-sources.json")
LATIN_FONT = "Helvetica Neue"
EAST_ASIA_FONT = "Microsoft YaHei"


def verify_reference(reference: Path, expected_sha256: str) -> None:
    digest = hashlib.sha256(reference.read_bytes()).hexdigest()
    if digest != expected_sha256:
        raise ValueError(
            f"retained reference changed: expected {expected_sha256}, got {digest}"
        )


def make_working_copy(reference: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(reference, output)


def set_east_asia_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def configure_styles(doc: Document) -> None:
    for style_name in ("normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def replace_preserving_run_format(paragraph: Paragraph, text: str) -> None:
    prototype_r_pr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        prototype_r_pr = deepcopy(paragraph.runs[0]._r.rPr)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if prototype_r_pr is not None:
        if run._r.rPr is not None:
            run._r.remove(run._r.rPr)
        run._r.insert(0, prototype_r_pr)
    set_east_asia_font(run)


def replace_cover_slots(doc: Document, meta: dict[str, str]) -> None:
    cover_image = doc.inline_shapes[0]
    aspect_ratio = cover_image.width / cover_image.height
    cover_image.height = Inches(4.85)
    cover_image.width = int(cover_image.height * aspect_ratio)
    title = next(p for p in doc.paragraphs if p.text == "Report title")
    cover_title = meta["title"].replace("与两年研究", "与\n两年研究")
    replace_preserving_run_format(title, cover_title)
    contents = next(p for p in doc.paragraphs if p.text == "Contents")
    replace_preserving_run_format(contents, "目录")
    cover_table = doc.tables[0]
    cover_subtitle = meta["subtitle"].replace("暑期强化计划", "\n暑期强化计划")
    replace_preserving_run_format(cover_table.cell(0, 0).paragraphs[0], cover_subtitle)
    metadata = f"编制：{meta['author']}\n{meta['date']}"
    metadata_cell = cover_table.cell(0, 2)
    replace_preserving_run_format(metadata_cell.paragraphs[0], metadata)
    for extra in list(metadata_cell.paragraphs[1:]):
        metadata_cell._tc.remove(extra._p)


def find_toc_sdt(doc: Document):
    for node in doc._element.body:
        if node.tag == qn("w:sdt") and "TOC" in "".join(node.itertext()):
            return node
    raise ValueError("template TOC content control was not found")


def clear_editable_template_body(doc: Document) -> None:
    body = doc._element.body
    toc_sdt = find_toc_sdt(doc)
    reached_toc = False
    for node in list(body):
        if node is toc_sdt:
            reached_toc = True
            continue
        if reached_toc and node.tag != qn("w:sectPr"):
            body.remove(node)


def clone_paragraph_properties(target: Paragraph, source_p_pr) -> None:
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    target._p.insert(0, deepcopy(source_p_pr))


def add_heading(
    doc: Document, text: str, level: int, page_break_before: bool = False
) -> Paragraph:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    return paragraph


def add_body(doc: Document, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    return paragraph


def add_label_body(doc: Document, label: str, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    label_run = paragraph.add_run(label)
    set_east_asia_font(label_run, bold=True)
    text_run = paragraph.add_run(text)
    set_east_asia_font(text_run)
    return paragraph


def add_bullet(doc: Document, text: str, bullet_p_pr) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    clone_paragraph_properties(paragraph, bullet_p_pr)
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    return paragraph


def add_numbered(doc: Document, text: str, numbered_p_pr) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    clone_paragraph_properties(paragraph, numbered_p_pr)
    run = paragraph.add_run(text)
    set_east_asia_font(run)
    return paragraph


def add_manual_numbered(
    doc: Document, index: int, text: str, numbered_p_pr, size: float | None = None
) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    clone_paragraph_properties(paragraph, numbered_p_pr)
    num_pr = paragraph._p.pPr.find(qn("w:numPr"))
    if num_pr is not None:
        paragraph._p.pPr.remove(num_pr)
    number_run = paragraph.add_run(f"{index}. ")
    set_east_asia_font(number_run, size=size, bold=True)
    text_run = paragraph.add_run(text)
    set_east_asia_font(text_run, size=size)
    return paragraph


def add_callout(doc: Document, label: str, text: str, callout_p_pr) -> Paragraph:
    paragraph = doc.add_paragraph(style="normal")
    clone_paragraph_properties(paragraph, callout_p_pr)
    label_run = paragraph.add_run(f"{label}。")
    set_east_asia_font(label_run, bold=True)
    text_run = paragraph.add_run(text)
    set_east_asia_font(text_run)
    return paragraph


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph(style="normal")
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, header: bool = False, size: float = 9.0) -> None:
    paragraph = cell.paragraphs[0]
    replace_preserving_run_format(paragraph, str(text))
    for run in paragraph.runs:
        set_east_asia_font(run, size=size, bold=header)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def set_table_widths(table: Table, widths: list[int]) -> None:
    grid_cols = table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_comparison_table(
    doc: Document,
    prototype_tbl,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
) -> Table:
    if len(headers) != 3 or any(len(row) != 3 for row in rows):
        raise ValueError("Design Report evidence tables must have exactly three columns")
    tbl_xml = deepcopy(prototype_tbl)
    tr_nodes = tbl_xml.findall(qn("w:tr"))
    desired = len(rows) + 1
    while len(tr_nodes) < desired:
        tbl_xml.append(deepcopy(tr_nodes[-1]))
        tr_nodes = tbl_xml.findall(qn("w:tr"))
    while len(tr_nodes) > desired:
        tbl_xml.remove(tr_nodes[-1])
        tr_nodes = tbl_xml.findall(qn("w:tr"))

    doc._element.body.insert(-1, tbl_xml)
    table = Table(tbl_xml, doc)
    for col, value in enumerate(headers):
        set_cell_text(table.cell(0, col), value, header=True)
    for row_index, row_values in enumerate(rows, start=1):
        for col, value in enumerate(row_values):
            set_cell_text(table.cell(row_index, col), value)
    set_table_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    return table


def add_section_intro(doc: Document, section_key: str, first: bool = False) -> None:
    section = SECTIONS[section_key]
    add_heading(doc, section["title"], 1, page_break_before=False)
    for paragraph in section.get("paragraphs", section.get("intro", [])):
        add_body(doc, paragraph)


def compose_executive_summary(doc: Document, bullet_p_pr, callout_p_pr) -> None:
    section = SECTIONS["executive_summary"]
    add_section_intro(doc, "executive_summary", first=True)
    add_heading(doc, "一目了然", 2)
    for item in section["at_a_glance"]:
        add_bullet(doc, item, bullet_p_pr)
    add_callout(doc, "核心判断", section["decision"], callout_p_pr)


def compose_generic_subsections(doc: Document, section_key: str) -> None:
    add_section_intro(doc, section_key)
    for subsection in SECTIONS[section_key]["subsections"]:
        add_heading(doc, subsection["title"], 2)
        for paragraph in subsection["paragraphs"]:
            add_body(doc, paragraph)


def compose_knowledge_architecture(doc: Document, bullet_p_pr) -> None:
    add_section_intro(doc, "knowledge_architecture")
    for layer in KNOWLEDGE_LAYERS:
        add_heading(doc, layer["name"], 2)
        add_label_body(doc, "核心问题：", "")
        for question in layer["core_questions"]:
            add_bullet(doc, question, bullet_p_pr)
        add_label_body(doc, "核心综述：", "；".join(layer["surveys"]))
        add_label_body(doc, "学习重点：", layer["what_to_learn"])
        add_label_body(doc, "工业映射：", layer["industrial_mapping"])
        add_label_body(doc, "阶段产出：", layer["output"])


def compose_costs(doc: Document, prototype_tbl, callout_p_pr) -> None:
    add_heading(doc, "个人研究原型成本与算力分工", 1, page_break_before=False)
    add_body(
        doc,
        "个人研究原型的主要成本不在购买大模型，而在数据、标注复核、批处理算力和可重复评测。以下金额为 2026 年个人研究级估算，不含个人时间，也不代表必须全部发生。",
    )
    add_heading(doc, "费用明细", 2)
    add_comparison_table(
        doc,
        prototype_tbl,
        ["项目", "参考费用", "控制方式"],
        [[name, cost, note] for name, cost, note in COST_ITEMS],
        [1900, 1500, 5720],
    )
    add_callout(
        doc,
        "预算建议",
        f"最低方案：{COST_SUMMARY['minimum']} 推荐方案：{COST_SUMMARY['recommended']} "
        f"个人上限：{COST_SUMMARY['upper_personal']} {COST_SUMMARY['control_rule']}",
        callout_p_pr,
    )
    add_heading(doc, "GPU 拿来做什么", 2)
    add_body(
        doc,
        "GPU 主要运行已经训练好的本地视觉分析模型，而不是默认训练新模型。感知结果应缓存，后续检索和 Agent 对比复用同一份结果，避免每轮实验重复消耗。",
    )
    add_comparison_table(
        doc,
        prototype_tbl,
        ["任务", "主要计算资源", "研究中的作用"],
        [[task, resource, role] for task, resource, role in GPU_ROLE_MATRIX],
        [2300, 1900, 4920],
    )


def week_title(index: int, week: dict[str, object]) -> str:
    start = str(week["start"])[5:].replace("-", " 月 ") + " 日"
    end = str(week["end"])[5:].replace("-", " 月 ") + " 日"
    return f"第 {index} 周｜{start}至 {end}｜{week['theme']}"


def compose_weekly_program(doc: Document, prototype_tbl, bullet_p_pr, callout_p_pr) -> None:
    add_section_intro(doc, "summer_program")
    summary_rows = [
        [f"第 {index} 周", str(week["theme"]), str(week["completion_gate"])]
        for index, week in enumerate(WEEKLY_PLAN, start=1)
    ]
    summary_rows.append(["8 月 31 日", str(FINALIZATION_DAY["theme"]), str(FINALIZATION_DAY["decision_rule"])])
    add_comparison_table(
        doc,
        prototype_tbl,
        ["阶段", "主题", "完成门槛"],
        summary_rows,
        [1100, 3000, 5020],
    )
    for index, week in enumerate(WEEKLY_PLAN, start=1):
        add_heading(doc, week_title(index, week), 2)
        add_label_body(doc, "建议投入：", str(week["hours"]))
        for label, field in (
            ("深读", "deep_reading"),
            ("选读", "selective_reading"),
            ("概念问题", "concept_questions"),
            ("实践", "practice"),
            ("交付物", "deliverables"),
        ):
            add_heading(doc, label, 3)
            for item in week[field]:
                add_bullet(doc, str(item), bullet_p_pr)
        add_label_body(doc, "选题漏斗动作：", str(week["topic_funnel_action"]))
        add_callout(doc, "本周完成门槛", str(week["completion_gate"]), callout_p_pr)
    add_heading(doc, "8 月 31 日｜选题冻结", 2)
    for action in FINALIZATION_DAY["actions"]:
        add_bullet(doc, str(action), bullet_p_pr)
    add_heading(doc, "最终交付", 3)
    for item in FINALIZATION_DAY["deliverables"]:
        add_bullet(doc, str(item), bullet_p_pr)
    add_callout(doc, "定题规则", str(FINALIZATION_DAY["decision_rule"]), callout_p_pr)


def compose_topic_funnel(doc: Document, prototype_tbl, bullet_p_pr, callout_p_pr) -> None:
    add_section_intro(doc, "topic_funnel")
    add_heading(doc, "四道硬门槛", 2)
    for gate in HARD_GATES:
        add_bullet(doc, gate, bullet_p_pr)
    add_callout(
        doc,
        "淘汰规则",
        "任何候选题只要有一道硬门槛不通过，就不进入加权评分；热点程度和个人兴趣不能抵消不可获得的数据或不可执行的评测。",
        callout_p_pr,
    )
    add_heading(doc, "100 分加权评分", 2)
    add_comparison_table(
        doc,
        prototype_tbl,
        ["指标", "权重", "评分证据"],
        [[label, str(weight), guide] for _, label, weight, guide in TOPIC_CRITERIA],
        [2300, 900, 5920],
    )
    add_heading(doc, "候选主题总览", 2)
    add_comparison_table(
        doc,
        prototype_tbl,
        ["主题", "核心研究问题", "主要风险"],
        [[f"{topic['id']} {topic['name']}", topic["research_question"], topic["main_risk"]] for topic in TOPIC_FAMILIES],
        [2100, 4200, 2820],
    )
    for topic in TOPIC_FAMILIES:
        add_heading(doc, f"{topic['id']}｜{topic['name']}", 2)
        for label, key in (
            ("研究问题：", "research_question"),
            ("可能创新：", "possible_innovation"),
            ("基线：", "baseline"),
            ("数据：", "data"),
            ("指标：", "metrics"),
            ("三个月原型：", "first_prototype"),
            ("论文收益：", "thesis_value"),
            ("工作收益：", "work_value"),
            ("主要风险：", "main_risk"),
        ):
            add_label_body(doc, label, str(topic[key]))


def compose_two_year_roadmap(doc: Document, prototype_tbl, bullet_p_pr) -> None:
    add_section_intro(doc, "two_year_roadmap")
    add_comparison_table(
        doc,
        prototype_tbl,
        ["阶段与时间", "核心目标", "可验收产出"],
        [[f"{phase}\n{dates}", goal, output] for phase, dates, goal, output in TWO_YEAR_PHASES],
        [1800, 2100, 5220],
    )
    add_heading(doc, "贯穿两年的可复用资产", 2)
    for item in (
        "版本化的视频事件数据模型与脱敏流程。",
        "含时间戳、证据和难度标记的评测集。",
        "本地视觉模型、检索器和大模型 API 的统一工具接口。",
        "可追踪模型、提示词、Token、GPU 时间、延迟和错误的运行日志。",
        "C++ 视频服务、Python 实验层与证据时间轴 UI。",
        "论文笔记、失败问题卡、实验记录和导师决策纪要。",
    ):
        add_bullet(doc, item, bullet_p_pr)


def compose_risks(doc: Document, prototype_tbl) -> None:
    add_section_intro(doc, "risks")
    add_comparison_table(
        doc,
        prototype_tbl,
        ["风险 / 等级", "预防措施", "停止或降级条件"],
        [[f"{name}\n{level}", mitigation, trigger] for name, level, mitigation, trigger in RISKS],
        [1800, 4000, 3320],
    )


def compose_recommendations(doc: Document, numbered_p_pr) -> None:
    add_section_intro(doc, "recommendations")
    for index, item in enumerate(SECTIONS["recommendations"]["items"], start=1):
        add_manual_numbered(doc, index, item, numbered_p_pr)


def compose_appendix(
    doc: Document,
    sources: list[dict[str, object]],
    prototype_tbl,
    bullet_p_pr,
    numbered_p_pr,
) -> None:
    add_section_intro(doc, "appendix")
    add_heading(doc, "A. 论文笔记模板", 2)
    for index, item in enumerate(APPENDICES["paper_note_template"], start=1):
        add_manual_numbered(doc, index, item, numbered_p_pr)
    add_heading(doc, "B. 周复盘清单", 2)
    for item in APPENDICES["weekly_review"]:
        add_bullet(doc, item, bullet_p_pr)
    add_heading(doc, "C. 候选题一页简报", 2)
    for index, item in enumerate(APPENDICES["candidate_brief"], start=1):
        add_manual_numbered(doc, index, item, numbered_p_pr)
    add_heading(doc, "D. 导师沟通简报", 2)
    for item in APPENDICES["advisor_brief"]:
        add_bullet(doc, item, bullet_p_pr)
    add_heading(doc, "E. 空白选题评分表", 2)
    add_comparison_table(
        doc,
        prototype_tbl,
        ["指标", "权重", "证据与得分"],
        [[label, str(weight), "证据：\n得分："] for _, label, weight, _ in TOPIC_CRITERIA],
        [2300, 900, 5920],
    )
    add_heading(doc, "F. 核心术语", 2)
    for term, definition in APPENDICES["glossary"]:
        add_label_body(doc, f"{term}：", definition)
    add_heading(doc, "G. 文献目录", 2)
    status_map = {
        "peer-reviewed": "同行评审",
        "accepted": "已接收",
        "preprint": "预印本",
        "status-unverified": "状态待核对",
    }
    mode_map = {"deep": "深读", "selective": "选读", "lookup": "查阅"}
    priority_order = {"core": 0, "targeted": 1, "reference": 2}
    ordered = sorted(
        sources,
        key=lambda item: (
            priority_order[str(item["priority"])],
            str(item["layer"]),
            str(item["id"]),
        ),
    )
    for index, source in enumerate(ordered, start=1):
        text = (
            f"{source['authors']}. {source['title']}. {source['year']}. "
            f"{status_map[str(source['status'])]}；{mode_map[str(source['reading_mode'])]}；"
            f"{source['id']}。{source['url']}"
        )
        add_manual_numbered(doc, index, text, numbered_p_pr, size=9.5)


def patch_header_text(path: Path, meta: dict[str, str]) -> None:
    temp = path.with_suffix(".header-patched.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temp, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("word/header") and item.filename.endswith(".xml"):
                text = data.decode("utf-8")
                text = text.replace("Report title", meta["header_title"])
                text = text.replace("Date", "2026 年 7 月")
                data = text.encode("utf-8")
            elif item.filename == "word/document.xml":
                text = data.decode("utf-8")
                text = re.sub(
                    r'(<w:instrText[^>]*>)[^<]*TOC[^<]*(</w:instrText>)',
                    r'\1 TOC \\o &quot;1-2&quot; \\h \\z \\u \2',
                    text,
                    count=1,
                )
                data = text.encode("utf-8")
            target.writestr(item, data)
    temp.replace(path)


def build_report(reference: Path, sources_path: Path, output: Path) -> None:
    verify_reference(reference, EXPECTED_SHA256)
    sources = json.loads(sources_path.read_text(encoding="utf-8"))
    make_working_copy(reference, output)
    doc = Document(output)
    configure_styles(doc)

    bullet_p_pr = deepcopy(doc.paragraphs[10]._p.pPr)
    callout_p_pr = deepcopy(doc.paragraphs[23]._p.pPr)
    numbered_p_pr = deepcopy(doc.paragraphs[29]._p.pPr)
    evidence_tbl = deepcopy(doc.tables[1]._tbl)

    replace_cover_slots(doc, REPORT_META)
    clear_editable_template_body(doc)
    add_page_break(doc)

    compose_executive_summary(doc, bullet_p_pr, callout_p_pr)
    compose_generic_subsections(doc, "background")
    compose_generic_subsections(doc, "key_findings")
    compose_knowledge_architecture(doc, bullet_p_pr)
    compose_costs(doc, evidence_tbl, callout_p_pr)
    compose_weekly_program(doc, evidence_tbl, bullet_p_pr, callout_p_pr)
    compose_generic_subsections(doc, "reading_method")
    compose_topic_funnel(doc, evidence_tbl, bullet_p_pr, callout_p_pr)
    compose_two_year_roadmap(doc, evidence_tbl, bullet_p_pr)
    compose_risks(doc, evidence_tbl)
    compose_recommendations(doc, numbered_p_pr)
    compose_appendix(doc, sources, evidence_tbl, bullet_p_pr, numbered_p_pr)

    doc.core_properties.title = REPORT_META["title"]
    doc.core_properties.subject = REPORT_META["subtitle"]
    doc.core_properties.author = REPORT_META["author"]
    doc.core_properties.keywords = "Agent, industrial video, survey, research planning"
    doc.save(output)
    patch_header_text(output, REPORT_META)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", type=Path, default=REFERENCE)
    parser.add_argument("--sources", type=Path, default=DEFAULT_SOURCES)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_report(args.reference, args.sources, args.output)
    print(f"Built {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
