from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("longvideobench-technical-report.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
LIGHT_FILL = "F2F4F7"
BORDER = "D9E2F3"
CJK_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size, bold=False, color=None, after=6, before=0, line=1.1):
    style.font.name = LATIN_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_widths(table, widths_in):
    widths = [int(w * 1440) for w in widths_in]
    total = sum(widths)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_paragraph_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:keepNext"))
    if value and existing is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and existing is not None:
        p_pr.remove(existing)


def add_footer_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("LongVideoBench Technical Report | ")
    set_run_font(run, size=9, color=MUTED)
    field_run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    field_run._r.append(fld_char_1)
    field_run._r.append(instr_text)
    field_run._r.append(fld_char_2)
    set_run_font(field_run, size=9, color=MUTED)


def add_para(doc, text="", style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text)
    set_run_font(r, size=16, bold=True, color=BLUE)
    set_paragraph_keep_with_next(p)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph(style="Subtopic")
    r = p.add_run(text)
    set_run_font(r, size=11.5, bold=True, color=DARK_BLUE)
    set_paragraph_keep_with_next(p)
    return p


def add_metric_table(doc):
    add_subheading(doc, "关键实验数字速览")
    rows = [
        [
            "实验维度",
            "原文配置/指标",
            "关键结果",
            "解读边界",
        ],
        [
            "数据规模",
            "3,763 个视频、6,678 道人工标注 MCQ；验证集 752 个视频/1,337 题，测试集 3,011 个视频/5,341 题；平均时长 473 秒。",
            "覆盖 Life、Movie、Knowledge、News 等主题，最长组为 (900s, 3600s]。",
            "这是评测基准规模，不等于可训练数据规模；测试集答案隐藏。",
        ],
        [
            "帧数扩展",
            "验证集按 max_frames 从 1 到 256 扫描，采样上限为 1 fps。",
            "GPT-4o 总准确率 41.7 -> 66.7，最长视频组 36.0 -> 60.9；Gemini-1.5-Pro 总准确率 38.6 -> 64.0，最长视频组 35.8 -> 58.6。",
            "原文证明更强模型受益于更多帧；没有证明所有模型都能靠堆帧提升。",
        ],
        [
            "开源模型缩放",
            "Idefics2、Mantis-Idefics2、Phi-3-Vision、Mantis-BakLLaVA 等按同一验证设置评估。",
            "Idefics2 在 16 帧为 49.7，64 帧降到 30.9；Mantis-Idefics2 在 16 帧为 47.0，64 帧降到 30.2；Phi-3-Vision 与 Mantis-BakLLaVA 在 64 帧设置下超过上下文长度。",
            "长上下文输入能力和实际有效利用能力不是一回事。",
        ],
        [
            "模态消融",
            "表 6 比较字幕-only、帧-only、帧+字幕三种输入。",
            "GPT-4o 为 44.6 / 60.6 / 66.7；Gemini-1.5-Pro 为 43.0 / 62.9 / 63.9；Idefics2 为 25.6 / 49.4 / 49.7。",
            "视觉帧是基础，字幕能补充信息；开源模型对字幕增益利用有限。",
        ],
        [
            "测试榜单",
            "表 7 报告长上下文、图像、多视频 LMM 的测试集结果。",
            "GPT-4o 测试总分 66.7，Gemini-1.5-Pro 64.4，Gemini-1.5-Flash 62.4，GPT-4-Turbo 60.7；最佳开源视频模型 PLLaVA-34B 为 53.5。",
            "排行榜说明能力差距，但不等价于部署成本、延迟或业务可用性。",
        ],
    ]

    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    set_table_widths(table, [1.0, 2.0, 2.4, 1.1])
    set_table_borders(table)
    for i, data_row in enumerate(rows):
        word_row = table.rows[i]
        for j, text in enumerate(data_row):
            cell = word_row.cells[j]
            if i == 0:
                shade_cell(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(text)
            set_run_font(r, size=9.2 if i else 9.5, bold=(i == 0), color=DARK_BLUE if i == 0 else None)
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.start_type = WD_SECTION_START.NEW_PAGE
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
add_footer_page_number(section)

styles = doc.styles
set_style_font(styles["Normal"], 11, after=6, before=0, line=1.1)
set_style_font(styles["Heading 1"], 16, bold=True, color=BLUE, before=16, after=8, line=1.1)
set_style_font(styles["Heading 2"], 13, bold=True, color=BLUE, before=12, after=6, line=1.1)
styles.add_style("Subtopic", 1)
set_style_font(styles["Subtopic"], 11.5, bold=True, color=DARK_BLUE, before=6, after=2, line=1.1)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.line_spacing = 1.05
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("LONGVIDEOBENCH 论文技术报告")
set_run_font(r, size=24, bold=True, color=BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
sr = subtitle.add_run("基于 arXiv:2407.15754v1 PDF 全文整理")
set_run_font(sr, size=11, color=MUTED)

metadata = [
    ("论文标题：", "LONG VIDEO BENCH: A Benchmark for Long-context Interleaved Video-Language Understanding"),
    ("作者：", "Haoning Wu, Dongxu Li, Bei Chen, Junnan Li"),
    ("版本信息：", "arXiv:2407.15754v1 [cs.CV], 22 Jul 2024；Preprint, Under review"),
    ("研究问题：", "如何评测大多模态模型在小时级、字幕与视频帧交织输入中的细节检索、时序关系理解和多模态推理能力。"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    lr = p.add_run(label)
    set_run_font(lr, bold=True, color=DARK_BLUE)
    vr = p.add_run(value)
    set_run_font(vr)

note = doc.add_paragraph()
note.paragraph_format.space_after = Pt(12)
nr = note.add_run("说明：以下分析只使用原文可见信息。凡原文未给出具体数值或未验证的落地结论，均按证据不足处理。")
set_run_font(nr, size=10.5, color=MUTED)

add_section_heading(doc, "1. 研究背景与行业痛点")
add_para(
    doc,
    "这篇论文关注的不是单个视频模型架构，而是一个更底层的评测问题：当大多模态模型（Large Multimodal Models, LMMs）已经开始支持 128K 乃至百万级上下文时，公开基准是否真的能衡量它们理解长视频的能力。原文指出，长上下文评测此前主要集中在纯文本输入，而长视频理解天然包含视觉帧、字幕、时序顺序和跨片段关系，是更贴近真实多模态应用的试验场。",
)
add_para(
    doc,
    "论文认为旧有视频问答基准存在明显的 single-frame bias：模型即使接收更多帧，指标也未必提升，说明问题可能只需要少量关键帧甚至单帧就能作答。这样的基准很难区分“模型真的能看懂长视频”和“模型碰巧看到了足够的静态线索”。对长视频检索、时序事件抽取和视频 Agent 来说，这个缺口会直接影响系统评估：如果测试题不迫使模型定位细节、比较片段和处理字幕，就无法暴露真实工程瓶颈。",
)
add_para(
    doc,
    "LONGVIDEOBENCH 试图补上的能力，是在小时级视频中同时考察细粒度检索和关系推理。基准包含 3,763 个带字幕的网页视频、6,678 道人工标注多选题，覆盖 Life、Movie、Knowledge、News 等主题，视频最长到 1 小时。论文把核心挑战概括为两件事：从很长的多模态输入里找回特定细节，以及围绕这些细节进行上下文关系推理。",
)
add_para(
    doc,
    "因此，这篇论文的行业价值更像“评测尺子”而不是“产品方案”。它没有声称提出一个可以直接部署的长视频理解系统，而是用更难被单帧捷径解决的问题，逼近真实长视频应用中的痛点：视频很长、信息分散、字幕和视觉互相补充、用户问题往往指向某个片段而不是整段视频摘要。",
)

add_section_heading(doc, "2. 核心创新亮点")
add_subheading(doc, "【理论创新】Referring Reasoning 把长视频理解拆成“定位 + 推理”")
add_para(
    doc,
    "论文提出 referring reasoning 作为基准的基础任务范式。每道题先包含一个 referring query，用自然语言指向视频中的一个或多个 referred context；随后问题主体要求模型围绕这个被指向的上下文作答。这个设计的关键不在于多选题形式本身，而在于它把长上下文评测从“整体摘要”转向“先找针，再理解针与周围内容的关系”。",
)
add_para(
    doc,
    "在理论拆分上，论文把问题分为 L1 Perception 和 L2 Relation。L1 主要考察单个片段中的对象、事件、属性等视觉感知；L2 要求模型关联多个片段，例如事件前后关系、对象追踪、场景顺序、属性变化等。17 个细粒度类别将 referring query 的来源（scene、event、object、subtitle）和目标答案类型（event、object、attribute、sequence 等）组合起来，使评测覆盖从静态识别到复杂时序理解的不同难度层级。",
)
add_subheading(doc, "【工程落地创新】数据构建和评估设置直接服务于长视频诊断")
add_para(
    doc,
    "工程侧的创新主要体现在数据和评估流程。论文从 119 个频道收集至少 720P 的网页视频，覆盖 10 个内容类别；对平台自带字幕进行筛选，对缺少字幕的视频使用 Whisper-V3-Large 转写；再用 Q-Align 过滤低质量视频，要求质量分数高于 0.25。最终人工筛到 3,763 个视频，并按 4 个递进时长组组织，从 8-15 秒到 15-60 分钟。",
)
add_para(
    doc,
    "标注环节也服务于“可定位性”。标注者必须完整观看视频，问题中要明确高亮 referred query，并给出 referred moment 的帧索引。每个视频经过 primary annotator、examiner、reviser 三个角色，原文称 20% 的标注被识别为有问题并修订。这个流程让基准不仅有答案标签，也保留了与时间位置相关的监督信号，适合后续分析 query depth 和长距离检索失败。",
)
add_para(
    doc,
    "需要强调的是，论文没有提出新的 LMM 训练方法，也没有证明某个模型结构是最优方案。它的创新边界是：通过任务定义、数据组织和系统化评估，构造一个更能暴露长视频多模态理解缺陷的公开基准。",
)

add_section_heading(doc, "3. 方法架构通俗拆解")
add_para(
    doc,
    "从输入看，LONGVIDEOBENCH 将视频帧和字幕组织成 temporally-aligned multimodal sequences。字幕块会按照字幕中间时间戳插入到前后相邻帧之间，让模型接收的不是“先一堆帧、再一段字幕”，而是更接近人类观看视频时的交织信息流。模型最后接收问题和选项，输出一个多选题字母。",
)
add_para(
    doc,
    "Referring query 可以理解为用户问题里的“定位锚点”。例如问题可能先描述某个画面、某个对象、某个事件，或者某句字幕出现时的片段；这些被锚定的画面和字幕就是 referred context。模型要先在长输入里找到这个上下文，再判断该上下文中的对象、颜色、动作、先后顺序或属性变化。",
)
add_para(
    doc,
    "L1 Perception 更像单点细节读取：在某个被指向片段中识别对象是否存在、对象是什么颜色、事件发生了什么。L2 Relation 则更接近时序理解和跨片段推理：同一对象在另一个场景是否出现、两个事件谁先谁后、多个场景的顺序是否正确、对象属性从一个片段到另一个片段如何变化。论文实验也显示，L2 Relation 整体比 L1 Perception 更困难。",
)
add_para(
    doc,
    "数据构建链路可以拆成四步：先按视频时长和主题类别收集候选视频；再保证视频有英文字幕，必要时用 Whisper-V3-Large 转写；然后通过 Q-Align 和人工筛选控制视频质量；最后由人工标注者创建 referring reasoning 问题、干扰选项、正确答案和 referred moment 帧索引。这里的关键不是“视频多”，而是每个问题都尽量绑定到可定位的细节证据。",
)
add_para(
    doc,
    "评估阶段，长上下文 LMM 以图像帧和文本字幕交织列表作为输入；图像 LMM 统一采样 8 帧；视频 LMM 按官方默认设置输入 16 或 32 帧。论文还专门分析 max_frames、输入模态和 referring query depth：前者看模型能否受益于更多帧，模态消融看字幕与视觉帧各自作用，query depth 则看被问片段离最终问题越远时模型是否更容易遗忘。",
)

add_section_heading(doc, "4. 实验配置与效果")
add_para(
    doc,
    "实验数据上，论文将 LONGVIDEOBENCH 分为验证集和测试集：验证集包含 752 个视频、1,337 道 MCQ；测试集包含 3,011 个视频、5,341 道 MCQ，测试答案隐藏以降低过拟合风险。表 1 对比了 MSVD-QA、MSRVTT-QA、ActivityNet-QA、MVBench、EgoSchema、MovieChat-1K 等常用视频基准，LONGVIDEOBENCH 的平均视频时长为 473 秒，且明确支持视频-语言交织输入。",
)
add_para(
    doc,
    "模型侧，论文评估了 GPT-4o、Gemini-1.5-Pro、GPT-4-Turbo、Gemini-1.5-Flash 等闭源长上下文 LMM，也包括 Idefics2、Phi-3-Vision-Instruct、Mantis 系列等开源长上下文模型，以及多图像 LMM 和视频 LMM。表 7 的 leaderboard 共列出 23 个模型条目。",
)
add_metric_table(doc)
add_para(
    doc,
    "最重要的实验结论是：对 GPT-4o、Gemini-1.5-Pro 等更强闭源模型，增加输入帧数确实能显著提高 LONGVIDEOBENCH 表现，尤其是超过 180 秒的视频。原文特别指出，在长于 180 秒的视频上，GPT-4o 和 Gemini-1.5-Pro 从 16 帧增加到 256 帧时可提升超过 10%。这说明该基准确实缓解了旧视频基准的 single-frame bias，模型需要处理更多帧才更可能答对。",
)
add_para(
    doc,
    "但开源模型呈现相反信号：Idefics2 和 Mantis-Idefics2 在超过 16 帧后不但没有稳定提升，64 帧时还明显下降；Phi-3-Vision 与 Mantis-BakLLaVA 在 64 帧设置下触发上下文长度限制。论文据此说明，能“接收长输入”和能“有效利用长输入”之间存在明显差距。",
)
add_para(
    doc,
    "输入模态消融也很清楚：只给字幕会显著低于给视频帧，说明视觉模态是基础；在帧上加入字幕通常会进一步提升，说明字幕对语义消歧和事件理解有帮助。不过，Idefics2 在帧-only 与帧+字幕之间仅从 49.4 到 49.7，提示部分开源模型尚不能很好整合长字幕和视觉线索。",
)
add_para(
    doc,
    "测试榜单上，GPT-4o 以 66.7 的测试总分最高，Gemini-1.5-Pro 为 64.4，Gemini-1.5-Flash 为 62.4，GPT-4-Turbo 为 60.7。开源模型中 PLLaVA-34B 以 53.5 居前；论文还指出，PLLaVA-13B 和 PLLaVA-34B 相比 PLLaVA-7B 分别提升 5.9% 和 14.3%，说明更强 LLM backbone 对综合视频理解有帮助。对于问题类型，原文强调 L2 Relation 比 L1 Perception 更难，SSS（Sequence of Scenes）是最困难类别之一；图 4 的 query depth 分析显示，referred moment 越靠近视频开头或中部，模型表现越差，但原文没有给出每个 depth 分桶的具体数值。",
)

add_section_heading(doc, "5. 研究局限与短板")
add_para(
    doc,
    "论文明确写出的局限有两点：当前 LONGVIDEOBENCH 只覆盖 vision 和 language 两种模态，没有纳入 audio；评估范围也没有包含超过 1 小时的视频。因此，对会议录音、现场采访、音乐/环境声强相关的视频，或者真正超长监控、直播回放、课程全集等场景，这个基准还不能完整代表真实任务。",
)
add_para(
    doc,
    "数据语言和来源也有边界。视频收集流程会移除没有字幕或非英文字幕的视频，缺字幕时用 Whisper-V3-Large 转写；这意味着论文没有验证中文、多语种字幕、强口音 ASR 噪声或无字幕视频下的表现。视频来自 119 个频道和 10 个主题类别，覆盖面比许多旧基准更广，但不能直接推断到所有行业视频分布。",
)
add_para(
    doc,
    "任务形式仍是多选问答。MCQ 有利于稳定评测和隐藏测试答案，但它不能完全替代开放式长视频检索、精确时间段定位（Temporal Grounding）、证据链引用或多轮 Agent 操作。也就是说，LONGVIDEOBENCH 能衡量模型是否能从候选答案里选对，但不能单独证明模型能在产品中可靠返回可审计的时间戳证据。",
)
add_para(
    doc,
    "人工标注虽然有三阶段质检，但原文也承认 examiner 和 reviser 识别并修订了 20% 有问题标注。这一数字说明 referring reasoning 问题本身标注难度较高，长视频中的片段描述、干扰项设计和 referred moment 标注仍可能存在残余歧义。论文提供了质量控制流程，但没有证明标注误差被完全消除。",
)
add_para(
    doc,
    "部署层面，实验结果本身也暴露了成本和上下文瓶颈：闭源模型在 256 帧下更强，但这通常意味着更高 token、图像输入和延迟成本；开源模型则可能在更多帧下退化或超过上下文长度。论文证明了 benchmark 难度和模型差距，但没有给出低成本长视频系统的工程解法。",
)
add_para(
    doc,
    "此外，数据集声明采用 CC BY-NC-SA 4.0 license，禁止商业使用。对于企业内部产品验证或商用插件评测，不能简单把数据集直接接入商业闭环，至少需要另建自有评测集或处理授权问题。",
)

add_section_heading(doc, "6. 落地启发（重点）")
add_subheading(doc, "可直接复用的模块或流程")
add_para(
    doc,
    "第一，可以直接复用“视频帧 + 字幕交织”的输入组织思路。对你的视频帧提取和长视频检索原型来说，字幕不应只作为独立文本索引存在，而可以按字幕 mid-timestamp 插入到相邻帧之间，形成时间对齐的多模态序列。这样做能让后续检索器或 LMM 在同一个上下文里看到画面和同期语义，减少“字幕说到了，但画面证据丢了”的问题。",
)
add_para(
    doc,
    "第二，referring query / referred context 可以作为时序事件抽取的标注规范。每个用户查询或评测问题都显式包含一个定位锚点，并记录 referred moment 的帧索引或时间戳。落到工程上，可以把一条样本拆成：查询锚点、候选时间段、视觉对象/事件、字幕片段、答案或事件标签。这比只保存整段视频摘要更适合训练和评估“先定位再推理”的模块。",
)
add_para(
    doc,
    "第三，17 个问题类别可以转成你自己的测试模板。S2O、S2A、T2O、T2A 适合检查帧提取和单帧细节识别；E3E、T3E、O3O 适合检查事件前后关系；SOS、TOS、SAA、TAA 适合检查对象追踪和属性变化；SSS 可以作为长视频时序排序的压力测试。这些模板能帮助 Agent 工作流从“能回答吗”细化到“是定位失败、视觉失败、字幕融合失败，还是时序推理失败”。",
)
add_para(
    doc,
    "第四，论文的 max_frames、模态消融和 query depth 分析可以直接转成评测仪表盘。一个可操作的原型评估方式是：固定同一批长视频问题，分别跑字幕-only、帧-only、帧+字幕；再按被问片段距离问题的远近分桶，统计准确率和召回时间段质量。这样能更快定位你的系统是检索召回弱，还是后端 LMM 对长上下文利用弱。",
)
add_subheading(doc, "需要二次验证后才能采用的思路")
add_para(
    doc,
    "论文显示 256 帧能显著提升 GPT-4o 和 Gemini-1.5-Pro，但这不应直接变成“所有长视频都喂 256 帧”的工程策略。你的场景需要重新测成本、延迟、视频类型和模型上下文预算。尤其是开源模型在表 5 中出现 64 帧退化或上下文超限，说明更现实的方案可能是分段索引、候选片段召回、再让强模型做局部推理，而不是全量堆帧。",
)
add_para(
    doc,
    "用 Whisper-V3-Large 生成字幕也值得复用，但需要在你的业务视频上单独评估 ASR 错误对检索和事件抽取的影响。论文保证所有视频有字幕，并强调字幕能提高理解效果；但它没有系统报告字幕错误率、中文语音、专业术语、嘈杂音频对 benchmark 的影响。对视频插件研发，字幕模块最好输出置信度、时间戳偏移和可回退的纯视觉路径。",
)
add_para(
    doc,
    "Q-Align 的质量过滤思路也可借鉴，但阈值不能照搬。论文用 Q-Align 过滤分数不高于 0.25 的低质量视频，这是为了构建评测集；你的原型如果处理真实用户视频，可能需要更细的质量维度，例如黑屏、镜头抖动、字幕遮挡、低光、画面重复、屏幕录制压缩伪影等。",
)
add_subheading(doc, "当前不适合落地或风险较高的部分")
add_para(
    doc,
    "不建议直接把 LONGVIDEOBENCH 分数当成视频产品能力指标。它是 MCQ benchmark，不覆盖开放式答案、证据引用、用户多轮追问、跨视频检索、音频线索和超过 1 小时的视频。更合理的做法是把它作为方法启发，然后构建与你的视频帧提取、时序事件抽取、长视频检索任务一致的内部评测集。",
)
add_para(
    doc,
    "也不建议直接商业复用该数据集。原文数据集声明使用 CC BY-NC-SA 4.0，非商业限制会影响企业插件或商业产品验证。可以复用的是任务范式、标注流程和评估维度；真正的商用评测数据最好来自自有视频、公开可商用素材或已授权数据。",
)
add_para(
    doc,
    "最后，不应把 Agent 工作流设计成“把长视频完整塞给一个大模型然后等待答案”。论文的 lost-in-context 信号很明确：query 越远或处于中部越容易出问题。更稳的 Agent 原型应拆成帧采样器、字幕对齐器、候选片段检索器、跨模态验证器和最终推理器；Agent 负责调度和自检，而不是替代所有底层检索与时序建模。",
)

doc.core_properties.title = "LONGVIDEOBENCH Technical Report"
doc.core_properties.author = ""
doc.core_properties.subject = "Chinese technical report for arXiv:2407.15754v1"
doc.core_properties.comments = "Generated from the local PDF and grounded in the original paper text."
doc.save(OUT)
print(OUT.resolve())
