from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("drvideo-technical-report.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
LIGHT_FILL = "F2F4F7"
BORDER = "D9E2F3"
CJK_FONT = "Microsoft YaHei"
LATIN_FONT = "Calibri"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size, bold=False, color=None, after=6, before=0, line=1.1):
    style.font.name = LATIN_FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_widths(table, widths_in):
    widths = [int(w * 1440) for w in widths_in]
    total = sum(widths)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def add_footer_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("DrVideo Technical Report | ")
    set_run_font(run, size=9, color=MUTED)
    field_run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    field_run._r.append(fld_char_1)
    field_run._r.append(instr_text)
    field_run._r.append(fld_char_2)
    set_run_font(field_run, size=9, color=MUTED)


def add_para(doc, text="", style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text)
    set_run_font(r, size=16, bold=True, color=BLUE)
    keep_with_next(p)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph(style="Subtopic")
    r = p.add_run(text)
    set_run_font(r, size=11.5, bold=True, color=DARK_BLUE)
    keep_with_next(p)
    return p


def add_metric_table(doc):
    add_subheading(doc, "关键实验数字速览")
    rows = [
        ["实验维度", "原文配置/指标", "关键结果", "解读边界"],
        [
            "数据集与设置",
            "EgoSchema：5000 个 3 分钟第一视角视频；MovieChat-1K：1000 个约 10 分钟影视视频；Video-MME long split：30-60 分钟，平均 44 分钟。",
            "EgoSchema 和 Video-MME 使用 accuracy；MovieChat-1K 使用 Gemini-Pro 辅助评估 accuracy 与 0-5 score。",
            "DrVideo 是 zero-shot / training-free 框架，不是新训练的端到端视频模型。",
        ],
        [
            "EgoSchema",
            "LaViLa 作为 captioner；比较 GPT-4 版本结果。",
            "DrVideo GPT-4：subset 66.4、fullset 61.0；VideoAgent GPT-4：60.2 / 54.1；VideoAgent [14]：62.8 / 60.2；LLoVi GPT-4 subset：61.2。",
            "DrVideo 在该表中领先 LLM-based 方法；提升来自文档检索和补充信息，不等于视觉编码能力本身提升。",
        ],
        [
            "MovieChat-1K",
            "LLaVA-NeXT 生成帧描述；global 与 breakpoint 两种模式。",
            "DrVideo：global Acc. 93.1、Score 4.41；breakpoint Acc. 56.4、Score 2.75。MovieChat+：71.2 / 3.51 与 49.6 / 2.62；VideoAgent*：65.4 / 3.12 与 31.6 / 2.05。",
            "Global 模式优势更大；breakpoint 仍只有 56.4，说明细节定位仍有空间。",
        ],
        [
            "Video-MME long split",
            "0.2 FPS 采样；w/o subs 与 w subs 两种设置；DeepSeek V2.5 作为 agent。",
            "DrVideo：w/o subs 51.7，w subs 71.7；subtitle-only 为 68.5。GPT-4o 为 65.3 / 72.1，Gemini 1.5 Pro 为 67.4 / 77.4，Gemini 1.5 Flash 为 61.1 / 68.8。",
            "加入字幕后 DrVideo 超过 Gemini 1.5 Flash、GPT-4o mini、GPT-4V、Claude 3.5 Sonnet，但没有超过 GPT-4o、Gemini 1.5 Pro 或 Qwen2-VL-72B。",
        ],
        [
            "组件消融",
            "EgoSchema subset，默认 GPT-3.5。",
            "从基础 57.4 加入 retrieval module 到 60.6，再加入 multi-stage agent interaction loop 到 62.6；CoT 从 62.2 提到 62.6。",
            "检索和 agent loop 是主增益；CoT 增益较小但可提升可解释性。",
        ],
        [
            "采样与模型选择",
            "Top-K、轮数、captioner、LLM、采样率消融。",
            "Top-K=5/10/20 为 62.6/61.4/60.6；采样率 1 FPS/0.5 FPS/0.25 FPS 为 61.6/62.6/58.8；LaViLa/LLaVA-NeXT/BLIP-2 为 62.6/61.2/59.6；GPT-4/GPT-3.5/DeepSeek/Mistral-8x7B 为 66.4/62.6/61.2/47.6。",
            "更多帧、更多轮、更多 key frame 不必然更好，噪声会干扰 LLM 判断。",
        ],
    ]

    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    set_table_widths(table, [1.0, 2.0, 2.4, 1.1])
    set_table_borders(table)
    for i, data_row in enumerate(rows):
        word_row = table.rows[i]
        for j, text in enumerate(data_row):
            cell = word_row.cells[j]
            if i == 0:
                shade_cell(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(text)
            set_run_font(r, size=9.0 if i else 9.5, bold=(i == 0), color=DARK_BLUE if i == 0 else None)
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.start_type = WD_SECTION_START.NEW_PAGE
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
add_footer_page_number(section)

styles = doc.styles
set_style_font(styles["Normal"], 11, after=6, before=0, line=1.1)
set_style_font(styles["Heading 1"], 16, bold=True, color=BLUE, before=16, after=8, line=1.1)
set_style_font(styles["Heading 2"], 13, bold=True, color=BLUE, before=12, after=6, line=1.1)
styles.add_style("Subtopic", 1)
set_style_font(styles["Subtopic"], 11.5, bold=True, color=DARK_BLUE, before=6, after=2, line=1.1)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.line_spacing = 1.05
r = title.add_run("DrVideo 论文技术报告")
set_run_font(r, size=24, bold=True, color=BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
sr = subtitle.add_run("基于 arXiv:2406.12846v2 PDF 全文整理")
set_run_font(sr, size=11, color=MUTED)

metadata = [
    ("论文标题：", "DrVideo: Document Retrieval Based Long Video Understanding"),
    ("作者：", "Ziyu Ma, Chenhui Gou, Hengcan Shi, Bin Sun, Shutao Li, Hamid Rezatofighi, Jianfei Cai"),
    ("机构：", "Hunan University；Monash University"),
    ("版本信息：", "arXiv:2406.12846v2 [cs.CV], 24 Nov 2024"),
    ("研究问题：", "如何在不重新训练大规模视频模型的前提下，通过文档检索、帧信息增强和 agent 迭代，提升长视频问答中的关键帧定位与长程推理能力。"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    lr = p.add_run(label)
    set_run_font(lr, bold=True, color=DARK_BLUE)
    vr = p.add_run(value)
    set_run_font(vr)

note = doc.add_paragraph()
note.paragraph_format.space_after = Pt(12)
nr = note.add_run("说明：以下分析只使用原文可见信息。凡原文没有实验验证的部署效果、业务价值或模型能力，均按证据不足处理。")
set_run_font(nr, size=10.5, color=MUTED)

add_section_heading(doc, "1. 研究背景与行业痛点")
add_para(
    doc,
    "这篇论文面向长视频理解（Long Video Understanding）里的两个核心难题：一是长视频帧数大幅增加后，模型很难定位与问题相关的关键视觉信息；二是即便找到了局部信息，模型还需要在较长时间跨度上进行推理。原文指出，已有视频理解方法多数集中在几十秒级短视频，对分钟级、十分钟级乃至更长视频的处理仍然不足。",
)
add_para(
    doc,
    "论文首先批评了直接把视频编码成视觉 token 再拼接给 LLM 的路线。以原文举例，OpenAI CLIP-L-14 每张图会输出 24x24 个 token，而 LLaVA-NeXT-Video 的最大长度只有 8192；这意味着完整长视频很难直接输入。为规避长度限制，许多方法使用均匀或随机稀疏采样，例如 PLLaVA 采 16 帧、IG-VLM 采 6 帧，但这种采样不看问题内容，容易漏掉真正关键帧。",
)
add_para(
    doc,
    "另一类方法把视频切成短片段，再转成文本 caption，让 LLM 做长文理解。LLoVi 是这一方向的代表，VideoAgent 进一步用 agent 和 CLIP 图文相似度粗到细地找关键帧。论文认为这些方法虽然缓解了长程推理问题，但仍有两类缺陷：其一，关键帧定位不够准，尤其在长视频中容易依赖 LLM 先验和粗粒度相似度；其二，即便找到了相关帧，caption 可能丢掉问题需要的细节，例如“女人照镜子”无法回答“她当时穿什么”。",
)
add_para(
    doc,
    "DrVideo 的研究动因就在这里：把长视频理解改造成长文档检索和理解问题。它试图让 LLM 先通过文本化视频获得全局理解，再围绕问题检索关键帧、补充缺失视觉信息，最后再回答。这一思路很贴近真实工程中的长视频分析：系统通常不能把所有帧都高成本精读，而需要先粗索引，再按问题做局部加密观察。",
)

add_section_heading(doc, "2. 核心创新亮点")
add_subheading(doc, "【理论创新】把长视频问答重新表述为长文档检索与理解")
add_para(
    doc,
    "DrVideo 的核心理论视角是 problem reformulation：不再把长视频直接看成超长视觉 token 序列，而是先把视频转成带 frame id 的视频文档（video document），再使用长文档检索与推理能力处理问题。这个转化使 LLM 能在语言空间里进行关键帧定位、信息充分性判断和长程推理，避免完全依赖视觉 token 上的端到端建模。",
)
add_para(
    doc,
    "它的关键区别不是“又做了一次 caption”，而是把 caption 变成可检索、可迭代增强的文档结构。初始文档只提供粗粒度帧描述；检索模块找出 Top-K 疑似关键帧；文档增强模块再对这些帧生成更详细、与问题相关的信息；agent loop 继续判断信息是否足够，并请求新的缺失帧或不同类型的增强信息。这使系统具备“先粗看，再追问，再补证据”的推理形态。",
)
add_subheading(doc, "【工程落地创新】训练无关的模块化长视频 Agent 框架")
add_para(
    doc,
    "工程上，DrVideo 是一个 training-free 框架，由五个模块组成：Video-Document Conversion、Document Retrieval、Document Augmentation、Multi-Stage Agent Interaction Loop、Answering Module。它不要求重新训练大型视频模型，而是组合现成 VLM、embedding model 和 LLM agent。原文称实验可在单张 NVIDIA RTX 4090 上复现，最低显存需求 24GB，并使用合理次数的 GPT 访问。",
)
add_para(
    doc,
    "Multi-stage agent interaction loop 是最有工程启发的部分。Planning agent 负责判断当前视频文档是否足以回答问题；如果不足，它写出缺什么信息；Interaction agent 再决定需要补哪些帧，以及每帧需要哪种信息：A 类是通用详细 caption，B 类是围绕问题做视觉问答（VQA）。补充信息进入文档后，planning agent 再次判断，直到信息足够或达到最大迭代轮数。",
)
add_para(
    doc,
    "需要边界清晰地说，论文没有提出新的视觉 backbone、视频 tokenizer 或训练数据配方。它的创新主要是系统设计：把长视频处理拆成文档化、检索、增强、计划、交互和回答几个可替换模块，并用实验说明这一组合在多个长视频 benchmark 上优于若干 LLM-based baseline。",
)

add_section_heading(doc, "3. 方法架构通俗拆解")
add_para(
    doc,
    "第一步是视频文档转换。给定长视频 V 和问题 Q，系统按采样率抽取帧或短 clip，并用 VLM 生成短描述。每个条目保留 frame id 和描述文本，形成 Docinit = {{1, SV1}, {2, SV2}, ..., {T, SVT}}。这里的 frame id 很关键，因为它让语言文档仍然保留视频时间顺序，后续检索到某句话时可以回到具体帧。",
)
add_para(
    doc,
    "第二步是文档检索。DrVideo 使用 OpenAI embedding model 把初始文档中的每个帧描述和问题文本编码成向量，然后按余弦相似度检索 Top-K 帧。论文默认 K 显著小于总帧数，例如 K=5 对比 T=90，因此只对少量候选帧做进一步增强，不会把文档长度暴力放大。",
)
add_para(
    doc,
    "第三步是文档增强。对检索出的关键帧，DrVideo 使用 LLaVA-NeXT 等 VLM 生成更详细的描述，或直接围绕问题对该帧做 VQA。原文给出的初始增强 prompt 大意是：如果问题中有事实错误，就精确描述图像；如果没有，则尝试回答问题。这一步是为了弥补普通短 caption 的信息丢失，例如服装颜色、物体位置、手部动作等细节。",
)
add_para(
    doc,
    "第四步是 agent 迭代。Planning agent 读取当前增强后的文档、问题和历史分析，判断信息是否足够。如果不够，它说明缺失点；Interaction agent 根据这些缺失点选择额外 N 个关键帧，并决定每个帧需要通用 caption 还是问题相关 VQA。新信息再追加回文档，形成下一轮输入。原文实验显示，DrVideo 在 EgoSchema 上第 2 轮达到峰值，过多轮数会因噪声信息增加而下降。",
)
add_para(
    doc,
    "最后是回答模块。系统把最终视频文档和问题交给 LLM，用 chain-of-thought 方式输出答案、置信度和解释。论文强调 CoT 不只是为了提升准确率，也有助于跟踪推理步骤和提升可解释性。不过，从消融结果看，CoT 的数值增益较小，主要提升来自检索和 agent loop。",
)

add_section_heading(doc, "4. 实验配置与效果")
add_para(
    doc,
    "论文在三个长视频 benchmark 上验证 DrVideo：EgoSchema、MovieChat-1K 和 Video-MME long split。EgoSchema 是 5000 个三分钟第一视角视频的多选问答，其中 500 题公开标签；MovieChat-1K 包含 1000 个约 10 分钟影视片段，区分 global 与 breakpoint 模式；Video-MME long split 视频长度为 30-60 分钟，平均 44 分钟。",
)
add_para(
    doc,
    "实现上，EgoSchema 使用 LaViLa 作为 captioner；MovieChat-1K 和 Video-MME 使用 LLaVA-NeXT 生成帧描述。采样率方面，EgoSchema 和 MovieChat-1K 为 0.5 FPS，Video-MME 为 0.2 FPS。比较实验中，EgoSchema 和 MovieChat-1K 使用 GPT-4 作为 agent；Video-MME 使用 DeepSeek V2.5；消融实验因 API 成本使用 GPT-3.5。所有 GPT-3.5、GPT-4 和 DeepSeek 实验 temperature 设为 0。",
)
add_metric_table(doc)
add_para(
    doc,
    "主结果显示，DrVideo 在 EgoSchema 上相较 LLoVi、VideoAgent、IG-VLM 等 LLM-based 方法有明确优势。尤其在使用相同视觉 captioner LaViLa 和 GPT-4 时，DrVideo subset accuracy 为 66.4，高于 VideoAgent 的 60.2 和 LLoVi GPT-4 的 61.2；fullset 为 61.0，高于 VideoAgent GPT-4 的 54.1，也略高于 VideoAgent [14] 的 60.2。",
)
add_para(
    doc,
    "在更长的 MovieChat-1K 上，DrVideo 的 global Acc. 达到 93.1、Score 4.41，明显高于 MovieChat+ 的 71.2 / 3.51 和 re-implemented VideoAgent 的 65.4 / 3.12。Breakpoint 模式更考验局部片段定位，DrVideo 为 56.4 / 2.75，高于 MovieChat+ 的 49.6 / 2.62 和 VideoAgent* 的 31.6 / 2.05，但绝对准确率仍说明细节定位没有完全解决。",
)
add_para(
    doc,
    "Video-MME long split 的结果更需要谨慎读。无字幕时 DrVideo 为 51.7，略高于 Claude 3.5 Sonnet 的 51.2，也高于 VideoAgent* 的 40.2 和 LLoVi* 的 45.4，但低于 GPT-4o、Gemini 1.5 Pro、Qwen2-VL-72B 等强模型。加入字幕后 DrVideo 达到 71.7，超过 Gemini 1.5 Flash、GPT-4o mini、GPT-4V、Claude 3.5 Sonnet，但仍低于 GPT-4o 的 72.1、Qwen2-VL-72B 的 74.3 和 Gemini 1.5 Pro 的 77.4。subtitle-only 已有 68.5，说明该 split 中字幕本身包含大量答案相关信息。",
)
add_para(
    doc,
    "消融实验支持论文的核心判断：只做基础文档理解时为 57.4，加入 retrieval module 后为 60.6，再加入 multi-stage agent interaction loop 后为 62.6；去掉 CoT 会从 62.6 降到 62.2。Top-K 从 5 增到 10、20 反而从 62.6 降到 61.4、60.6；采样率从 0.5 FPS 提到 1 FPS 也从 62.6 降到 61.6，降到 0.25 FPS 则为 58.8。结论很明确：信息太少会漏关键帧，信息太多又会引入噪声。",
)

add_section_heading(doc, "5. 研究局限与短板")
add_para(
    doc,
    "论文明确写出的第一项局限是上下文长度瓶颈。DrVideo 能处理的最长视频取决于 LLM 的最大 token length；如果视频远长于现有 benchmark，例如 10 小时级别，初始文档、增强描述、历史分析和字幕都可能把上下文撑爆。这说明 DrVideo 不是无限长视频方案，仍需要分层索引、摘要压缩或外部记忆。",
)
add_para(
    doc,
    "第二项原文局限是信息生成与噪声控制仍有改进空间。消融实验已经说明，Top-K 增大、迭代轮数过多、采样率过高都可能引入冗余和无关信息，干扰 LLM 判断。换句话说，DrVideo 的关键不是“多看几帧”，而是“选对要细看的帧”。这也是实际长视频系统最难稳定调参的部分。",
)
add_para(
    doc,
    "补充材料中的 failure case 进一步揭示了 VLM 依赖问题：当 VLM 无法准确描述视频内容时，DrVideo 会在错误文档上做后续 planning、interaction 和 answering，最终做出错误判断。论文明确指出 DrVideo heavily relies on both LLMs and VLMs，因此它的上限很大程度受基础模型能力影响。",
)
add_para(
    doc,
    "此外，Video-MME 的 subtitle-only 结果达到 68.5，而 DrVideo 加字幕为 71.7，视觉增强带来的额外增益是 3.2。这说明在某些长视频问答中，字幕已经承担了大量信息来源；如果落地场景没有高质量字幕，或者字幕与画面强错位，论文结果不能直接外推。",
)
add_para(
    doc,
    "评估形态也有边界。EgoSchema 和 Video-MME 是多选题，MovieChat-1K 依赖 GPT-assisted evaluation。论文证明了 benchmark 分数提升，但没有系统评估开放式时间戳检索、证据引用、交互式用户追问、实时处理延迟、API 成本、隐私约束和在线部署鲁棒性。因此，把 DrVideo 当作产品架构时仍需补大量工程验证。",
)

add_section_heading(doc, "6. 落地启发（重点）")
add_subheading(doc, "可直接复用的模块或流程")
add_para(
    doc,
    "对视频帧提取来说，最值得复用的是“帧 ID + 时间顺序 + 短描述”的视频文档结构。你的原型可以把每个采样帧保存为 frame_id、timestamp、brief_caption、source_clip、embedding 等字段，让后续检索和 agent 都能回到具体帧，而不是只保留一段整体摘要。",
)
add_para(
    doc,
    "对长视频检索来说，可以直接复用 DrVideo 的两阶段策略：先用低成本 caption 建立全局文本索引，再根据用户问题做 Top-K 检索，随后只对候选帧调用更贵的 VLM/VQA 增强。这个流程比全帧精读更省成本，也比固定间隔采样更贴近问题。",
)
add_para(
    doc,
    "对时序事件抽取来说，DrVideo 的 document augmentation 可以改造成事件补证模块。初始 caption 负责粗粒度事件候选，interaction agent 根据问题或缺失槽位请求“动作细节”“对象状态”“人物关系”“前后变化”等不同增强类型。这样可以把事件抽取拆成可审计的多步流程。",
)
add_para(
    doc,
    "对 Agent 工作流原型来说，planning agent / interaction agent / answering agent 的分工可以直接借鉴。Planning agent 不急着回答，而是先判断证据是否足够；Interaction agent 只负责提出要补哪些帧和什么类型信息；Answering agent 基于最终文档回答。这个分工能减少一个大 prompt 同时做所有事情带来的混乱。",
)
add_subheading(doc, "需要二次验证后才能采用的思路")
add_para(
    doc,
    "采样率不能照搬。论文在 EgoSchema 上 0.5 FPS 最好，Video-MME 使用 0.2 FPS，但这只是特定 benchmark 和任务下的结果。你的场景如果是视频插件、教程、会议、监控或影视片段，最佳采样率会不同。更稳妥的方向是自适应采样：先低频扫全局，再对事件密集或检索命中的片段加密抽帧。",
)
add_para(
    doc,
    "Top-K 和 agent 轮数也要二次验证。论文中 Top-K=5 优于 10/20，DrVideo 在 2 轮达到峰值，说明更多信息会带来噪声。但不同视频长度、caption 质量、问题类型和 LLM 上下文能力都会改变最佳点。工程上应把 K、轮数、每轮新增帧数和增强类型做成可观测参数，而不是写死。",
)
add_para(
    doc,
    "字幕融合值得采用，但必须单独评估。Video-MME 上字幕贡献很大，然而你的数据如果包含中文口语、行业术语、ASR 错误、多人重叠讲话或字幕缺失，subtitle-only 的优势可能消失。建议把字幕置信度、时间戳对齐质量和纯视觉回退路径纳入评估。",
)
add_subheading(doc, "当前不适合落地或风险较高的部分")
add_para(
    doc,
    "不建议直接把整段长视频全部转成一个无限增长的文档交给 LLM。论文已经指出 token length 是瓶颈；增强信息、历史分析和字幕越积越多，10 小时级视频会很快失控。更可落地的结构是分层文档：clip-level 索引、segment-level 摘要、frame-level 证据，回答时只装载相关上下文。",
)
add_para(
    doc,
    "也不建议让 VLM caption 成为唯一视觉证据。failure case 表明，一旦 caption 错了，后续 agent 会在错误文本上越推越远。工程原型应保留帧图像引用和时间戳，在关键答案前做二次视觉核验，尤其是颜色、数量、文字、手部动作、空间关系等容易被 caption 模糊化的细节。",
)
add_para(
    doc,
    "最后，DrVideo 依赖外部 LLM/VLM/API 的能力和成本。原文证明其 researcher-friendly，但没有给出生产环境成本、延迟或隐私方案。对你的视频插件研发来说，最好先把它作为离线分析和评测原型，用来验证“检索-补证-回答”的链路，再逐步替换成本更可控的本地模型或缓存策略。",
)

doc.core_properties.title = "DrVideo Technical Report"
doc.core_properties.author = ""
doc.core_properties.subject = "Chinese technical report for arXiv:2406.12846v2"
doc.core_properties.comments = "Generated from the local PDF and grounded in the original paper text."
doc.save(OUT)
print(OUT.resolve())
